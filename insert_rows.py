from docx import Document

# 打开Word文档
doc = Document('doc.docx')

# 获取表格
table = doc.tables[0]  # 假设要操作的表格是第一个表格

# 定义目标参数和要插入的行数
target_parameter1 = '距离（m）'
target_parameter2 = '备注信息'
num_insert_rows1 = 4  # 示例值，根据实际需求修改

num_insert_rows2 = num_insert_rows1

# 定义插入行的函数
def insert_rows_below(table, row_index, num_insert_rows):
    # 获取目标行
    target_row = table.rows[row_index]

    # 获取目标行的下一行索引
    next_row_index = row_index + 1

    # 在目标行下方插入指定数量的行
    for _ in range(num_insert_rows):
        # 创建新行
        new_row = table.add_row()

        # 将新行移动到目标行的下方
        table._tbl.remove(new_row._tr)
        table._tbl.insert(next_row_index , new_row._tr)

# 遍历表格行
for i,row in enumerate(table.rows):
    # 检查当前行是否包含目标参数1
    # if any(cell.text.strip() == target_parameter1 for cell in row.cells):
    #     # 获取当前行的索引
    #     row_index = i
    for cell in row.cells:
        if cell.text.strip() == target_parameter1:
            row_index = i
            print(cell.text)
            print(i)

        # 在目标行下方插入指定数量的行
            insert_rows_below(table, row_index +2 , num_insert_rows1)

# 再次遍历表格行
for i, row in enumerate(table.rows):
    # 检查当前行是否包含目标参数2
    # if any(cell.text.strip() == target_parameter2 for cell in row.cells):
    for cell in row.cells:
        if cell.text.strip() == target_parameter2:
        # 获取当前行的索引
            row_index = i
            print(cell.text)
            print(i)

            # 计算要插入的行数
            # 当num_insert_rows2为偶数时，row_to_insert等于num_insert_rows2
            # 当num_insert_rows2为奇数时，row_to_insert等于num_insert_rows2 + 1
            if num_insert_rows2 % 2 == 0:
                 rows_to_insert = num_insert_rows2
            else:
                 rows_to_insert = num_insert_rows2 + 1

            # 在目标行下方插入指定数量的行
            insert_rows_below(table, row_index + 2 , rows_to_insert)

# 单元格进行合并
def merge(start_row_index,start_column_index,end_row_index,end_column_index):
    table.cell(start_row_index, start_column_index).merge(table.cell(end_row_index, end_column_index))

merge(12,0,12,4)
merge(12,5,12,7)
merge(13,0,13,4)
merge(13,5,13,7)
merge(14,0,14,4)
merge(14,5,14,7)
merge(15,0,15,4)
merge(15,5,15,7)

merge(7,4,7,5)
merge(8,4,8,5)
merge(9,4,9,5)
merge(10,4,10,5)

# 保存并关闭文档
doc.save('new_doc.docx')
